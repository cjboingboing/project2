"""
Downloads the SA2-level ABS data and builds a postcode -> SA2 correspondence.

The consumer table only gives us a postcode, but postcodes are a postal
convenience, not a statistical geography: the ABS publishes most of its useful
data (SEIFA, and anything needing a boundary to draw) on SA2s instead. This
script builds the bridge between the two, population-weighted at Mesh Block
level, so consumers can be placed in an SA2 and mapped.

Outputs (data/tables/processed/):
    abs_census_sa2              one row per SA2: Census G01/G02 + SEIFA indexes
    abs_poa_sa2_correspondence  every postcode x SA2 overlap, with its weight
    abs_poa_sa2_primary         one row per postcode: its dominant SA2

Also writes data/tables/external/SA2_2021_AUST_simplified.geojson for mapping.

Run with:
    python scripts/download_external_sa2.py
"""

import csv
import shutil
import urllib.request
import zipfile
from pathlib import Path

import openpyxl
from pyspark.sql import SparkSession, Window, functions as F

#Paths

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
DATA_DIR = PROJECT_ROOT / "data" / "tables"
EXTERNAL_DIR = DATA_DIR / "external"
OUTPUT_DIR = DATA_DIR / "processed"

ASGS_BASE = (
    "https://www.abs.gov.au/statistics/standards/"
    "australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/"
    "access-and-downloads/"
)

CENSUS_SA2_URL = (
    "https://www.abs.gov.au/census/find-census-data/datapacks/download/"
    "2021_GCP_SA2_for_AUS_short-header.zip"
)
SEIFA_SA2_URL = (
    "https://www.abs.gov.au/statistics/people/people-and-communities/"
    "socio-economic-indexes-areas-seifa-australia/2021/"
    "Statistical%20Area%20Level%202%2C%20Indexes%2C%20SEIFA%202021.xlsx"
)
MB_COUNTS_URL = (
    "https://www.abs.gov.au/census/guide-census-data/mesh-block-counts/2021/"
    "Mesh%20Block%20Counts%2C%202021.xlsx"
)
MB_ALLOCATION_URL = ASGS_BASE + "allocation-files/MB_2021_AUST.xlsx"
POA_ALLOCATION_URL = ASGS_BASE + "allocation-files/POA_2021_AUST.xlsx"
SA2_BOUNDARY_URL = ASGS_BASE + "digital-boundary-files/SA2_2021_AUST_SHP_GDA2020.zip"

ZIP_TABLE_DIR = "2021 Census GCP Statistical Area 2 for AUS/"
G01_CSV = "2021Census_G01_AUST_SA2.csv"
G02_CSV = "2021Census_G02_AUST_SA2.csv"

# raw ABS column -> our column. Kept deliberately in step with the postcode-level
# table in download_external.py so the two can be compared like for like.
G01_COLUMNS = {
    "Tot_P_P": "abs_population",
}
G02_COLUMNS = {
    "Median_age_persons": "abs_median_age",
    "Median_tot_prsnl_inc_weekly": "abs_median_personal_income_weekly",
    "Median_tot_hhd_inc_weekly": "abs_median_household_income_weekly",
    "Median_rent_weekly": "abs_median_rent_weekly",
    "Median_mortgage_repay_monthly": "abs_median_mortgage_repay_monthly",
}

# Table 1 of the SEIFA workbook, by position: the header is two rows deep
# (index name on one row, Score/Decile on the next), so it is easier to take
# the columns by index than to try and parse it.
SEIFA_COLUMNS = [
    "sa2_code",
    "sa2_name",
    "seifa_irsd_score",
    "seifa_irsd_decile",
    "seifa_irsad_score",
    "seifa_irsad_decile",
    "seifa_ier_score",
    "seifa_ier_decile",
    "seifa_ieo_score",
    "seifa_ieo_decile",
]

# ABS serves 403s to the default urllib user agent.
USER_AGENT = "Mozilla/5.0 (MAST30034 student project)"


def download(url, dest):
    """Download a file, skipped if already on disk."""
    if dest.exists():
        print(f"Already downloaded, skipping: {dest.name}")
        return dest

    dest.parent.mkdir(parents=True, exist_ok=True)
    partial = dest.with_name(dest.name + ".part")
    print(f"Downloading {dest.name} ...")
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request) as response, open(partial, "wb") as f:
        shutil.copyfileobj(response, f)
    # rename only once complete
    partial.rename(dest)
    return dest


def extract_census_tables(zip_path, out_dir=EXTERNAL_DIR):
    """Extract G01 and G02 from the SA2 DataPack."""
    paths = {}
    with zipfile.ZipFile(zip_path) as z:
        for name in (G01_CSV, G02_CSV):
            target = out_dir / name
            with z.open(ZIP_TABLE_DIR + name) as src, open(target, "wb") as dst:
                shutil.copyfileobj(src, dst)
            paths[name] = target
    return paths


def sheet_to_csv(xlsx_path, csv_path, header_marker, columns, sheets=None):
    """Stream rows out of an ABS workbook into a flat CSV.

    ABS workbooks carry several rows of title blurb above the real header and
    split large tables over numbered sheets, so rows are taken by position
    after the row whose first cell is `header_marker`. The CSV is cached
    because openpyxl is slow over these (~370k rows) and nothing here changes
    between runs.
    """
    if csv_path.exists():
        print(f"Already converted, skipping: {csv_path.name}")
        return csv_path

    print(f"Converting {xlsx_path.name} -> {csv_path.name} ...")
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    n_rows = 0
    try:
        with open(csv_path, "w", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            for sheet in sheets or workbook.sheetnames:
                found_header = False
                for row in workbook[sheet].iter_rows(values_only=True):
                    if not found_header:
                        found_header = row and row[0] == header_marker
                        continue
                    if row[0] is None:  # trailing footnotes
                        break
                    writer.writerow(row[: len(columns)])
                    n_rows += 1
    finally:
        workbook.close()

    print(f"  {n_rows:,} rows")
    return csv_path


def seifa_sheets(xlsx_path):
    """The one summary sheet of the SEIFA workbook that holds all four indexes."""
    return ["Table 1"]


def mb_counts_sheets(xlsx_path):
    """Every state sheet of the Mesh Block counts workbook ("Table 1", "Table 1.1", ...)."""
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
    try:
        return [s for s in workbook.sheetnames if s.startswith("Table ")]
    finally:
        workbook.close()


def build_census_table(spark, g01_path, g02_path, seifa_path):
    """One row per SA2: Census G01/G02 medians plus the four SEIFA indexes."""
    g01 = spark.read.option("header", "true").csv(str(g01_path))
    g02 = spark.read.option("header", "true").csv(str(g02_path))

    g01 = g01.select(
        F.col("SA2_CODE_2021").alias("sa2_code"),
        *[F.col(raw).cast("int").alias(new) for raw, new in G01_COLUMNS.items()],
    )
    g02 = g02.select(
        F.col("SA2_CODE_2021").alias("sa2_code"),
        *[F.col(raw).cast("int").alias(new) for raw, new in G02_COLUMNS.items()],
    )
    census = g01.join(g02, on="sa2_code", how="inner")

    n_g01, n_g02, n_joined = g01.count(), g02.count(), census.count()
    assert n_g01 == n_g02 == n_joined, f"G01={n_g01}, G02={n_g02}, joined={n_joined}"

    # ABS uses 0 for medians it cannot calculate (too few people to publish)
    for col in G02_COLUMNS.values():
        census = census.withColumn(col, F.when(F.col(col) == 0, F.lit(None)).otherwise(F.col(col)))

    seifa = spark.read.option("header", "true").csv(str(seifa_path))
    # SEIFA writes "-" where an index could not be built (SA2s with too few
    # people), so these have to be try_cast, not cast, or the whole job dies.
    seifa = seifa.select(
        F.col("sa2_code").cast("string").alias("sa2_code"),
        F.trim("sa2_name").alias("sa2_name"),
        *[
            F.col(c).try_cast("double" if c.endswith("score") else "int").alias(c)
            for c in SEIFA_COLUMNS[2:]
        ],
    )

    # SEIFA has no row for SA2s with no usual residents (airports, industrial
    # parks, "Migratory - Offshore - Shipping"), so this has to stay a left join.
    census = census.join(seifa, on="sa2_code", how="left")
    return census


def build_correspondence(spark, mb_alloc_path, poa_alloc_path, mb_counts_path):
    """Population-weighted postcode -> SA2 correspondence, built at Mesh Block level.

    Mesh Blocks are the smallest ABS geography and nest inside both postal areas
    and SA2s, so overlaying the two allocation files gives every postcode/SA2
    pair. Each pair is then weighted by the share of the postcode's Census
    population living in it.
    """
    mb_sa2 = spark.read.option("header", "true").csv(str(mb_alloc_path)).select(
        F.col("MB_CODE_2021").alias("mb_code"),
        F.col("SA2_CODE_2021").alias("sa2_code"),
        F.trim("SA2_NAME_2021").alias("sa2_name"),
        F.trim("STATE_NAME_2021").alias("state_name"),
    )
    mb_poa = spark.read.option("header", "true").csv(str(poa_alloc_path)).select(
        F.col("MB_CODE_2021").alias("mb_code"),
        F.lpad(F.trim("POA_CODE_2021"), 4, "0").alias("postcode"),
    )
    mb_counts = spark.read.option("header", "true").csv(str(mb_counts_path)).select(
        F.col("MB_CODE_2021").alias("mb_code"),
        F.col("Person").cast("int").alias("persons"),
        F.col("Dwelling").cast("int").alias("dwellings"),
    )

    mesh_blocks = mb_poa.join(mb_sa2, on="mb_code", how="inner").join(
        mb_counts, on="mb_code", how="left"
    )

    n_poa, n_overlaid = mb_poa.count(), mesh_blocks.count()
    assert n_poa == n_overlaid, f"POA={n_poa}, overlaid={n_overlaid} (mesh blocks lost or duplicated)"

    pairs = mesh_blocks.groupBy("postcode", "sa2_code", "sa2_name", "state_name").agg(
        F.sum(F.coalesce("persons", F.lit(0))).alias("persons"),
        F.sum(F.coalesce("dwellings", F.lit(0))).alias("dwellings"),
        F.count("*").alias("mesh_blocks"),
    )

    # Population is the right weight, but a handful of postcodes are entirely
    # non-residential (industrial estates, airports, CBD blocks) and have no
    # residents to weight by, so fall back to dwellings, then to a plain mesh
    # block count, rather than dropping the postcode.
    by_postcode = Window.partitionBy("postcode")
    pairs = (
        pairs.withColumn("postcode_persons", F.sum("persons").over(by_postcode))
        .withColumn("postcode_dwellings", F.sum("dwellings").over(by_postcode))
        .withColumn("postcode_mesh_blocks", F.sum("mesh_blocks").over(by_postcode))
    )
    pairs = pairs.withColumn(
        "weight_basis",
        F.when(F.col("postcode_persons") > 0, F.lit("persons"))
        .when(F.col("postcode_dwellings") > 0, F.lit("dwellings"))
        .otherwise(F.lit("mesh_blocks")),
    ).withColumn(
        "weight",
        F.when(F.col("weight_basis") == "persons", F.col("persons") / F.col("postcode_persons"))
        .when(F.col("weight_basis") == "dwellings", F.col("dwellings") / F.col("postcode_dwellings"))
        .otherwise(F.col("mesh_blocks") / F.col("postcode_mesh_blocks")),
    )

    # Ties (two SA2s with identical weight) are broken on SA2 code so that the
    # primary SA2 is stable between runs.
    ranked = Window.partitionBy("postcode").orderBy(F.desc("weight"), F.asc("sa2_code"))
    pairs = pairs.withColumn("weight_rank", F.row_number().over(ranked))

    correspondence = pairs.select(
        "postcode", "sa2_code", "sa2_name", "state_name",
        "persons", "dwellings", "mesh_blocks",
        "weight_basis", "weight", "weight_rank",
    )

    primary = (
        correspondence.filter(F.col("weight_rank") == 1)
        .withColumn("sa2_share", F.col("weight"))
        .join(
            correspondence.groupBy("postcode").agg(F.count("*").alias("n_sa2s")),
            on="postcode",
            how="inner",
        )
        .select(
            "postcode", "sa2_code", "sa2_name", "state_name",
            "sa2_share", "n_sa2s", "weight_basis",
        )
    )

    return correspondence, primary


def write_boundaries(shp_zip_path, out_path, tolerance=0.001):
    """Write SA2 boundaries as GeoJSON, simplified enough to draw quickly.

    The full shapefile is ~70MB of coastline detail that no choropleth at this
    scale can show; a ~100m tolerance keeps the shapes recognisable and makes
    the file small enough to load in a notebook.
    """
    if out_path.exists():
        print(f"Already built, skipping: {out_path.name}")
        return out_path

    import geopandas as gpd  # only needed here, keeps the ETL importable without GDAL

    print("Building simplified SA2 boundaries ...")
    boundaries = gpd.read_file(f"zip://{shp_zip_path}")
    boundaries = boundaries[["SA2_CODE21", "SA2_NAME21", "STE_NAME21", "geometry"]].rename(
        columns={"SA2_CODE21": "sa2_code", "SA2_NAME21": "sa2_name", "STE_NAME21": "state_name"}
    )
    # SA2s with no geometry are the non-spatial ones (no usual address, migratory)
    n_before = len(boundaries)
    boundaries = boundaries[boundaries.geometry.notna()]
    print(f"  {len(boundaries):,} SA2s with geometry ({n_before - len(boundaries)} non-spatial dropped)")

    boundaries = boundaries.to_crs(epsg=4326)
    boundaries["geometry"] = boundaries.geometry.simplify(tolerance, preserve_topology=True)
    boundaries.to_file(out_path, driver="GeoJSON")
    print(f"  Written to {out_path} ({out_path.stat().st_size / 1e6:.1f}MB)")
    return out_path


def report_correspondence(correspondence, primary):
    """Print how cleanly postcodes sit inside SA2s."""
    width = 58
    n_postcodes = primary.count()

    print(f"{' postcode -> SA2 correspondence ':-^{width}}")
    print(f"{'postcodes':<38}{n_postcodes:>{width - 38},}")
    print(f"{'postcode x SA2 pairs':<38}{correspondence.count():>{width - 38},}")
    print(f"{'SA2s covered':<38}{correspondence.select('sa2_code').distinct().count():>{width - 38},}")

    print(f"{' postcodes split across SA2s ':-^{width}}")
    split = primary.select(
        F.sum((F.col("n_sa2s") == 1).cast("int")).alias("one"),
        F.sum((F.col("n_sa2s").between(2, 5)).cast("int")).alias("few"),
        F.sum((F.col("n_sa2s") > 5).cast("int")).alias("many"),
        F.max("n_sa2s").alias("most"),
    ).first()
    print(f"  {'wholly inside one SA2':<36}{split['one']:>{width - 38},}  ({split['one'] / n_postcodes:.1%})")
    print(f"  {'spread over 2-5 SA2s':<36}{split['few']:>{width - 38},}  ({split['few'] / n_postcodes:.1%})")
    print(f"  {'spread over 6+ SA2s':<36}{split['many']:>{width - 38},}  ({split['many'] / n_postcodes:.1%})")
    print(f"  {'most SA2s for one postcode':<36}{split['most']:>{width - 38},}")

    print(f"{' confidence in the primary SA2 ':-^{width}}")
    share = primary.select(
        F.sum((F.col("sa2_share") >= 0.9).cast("int")).alias("strong"),
        F.sum((F.col("sa2_share").between(0.5, 0.9)).cast("int")).alias("fair"),
        F.sum((F.col("sa2_share") < 0.5).cast("int")).alias("weak"),
        F.mean("sa2_share").alias("mean_share"),
    ).first()
    print(f"  {'primary SA2 holds 90%+ of postcode':<36}{share['strong']:>{width - 38},}  ({share['strong'] / n_postcodes:.1%})")
    print(f"  {'holds 50-90%':<36}{share['fair']:>{width - 38},}  ({share['fair'] / n_postcodes:.1%})")
    print(f"  {'holds under 50%':<36}{share['weak']:>{width - 38},}  ({share['weak'] / n_postcodes:.1%})")
    print(f"  {'mean share':<36}{share['mean_share']:>{width - 38}.3f}")

    print(f"{' what the weight was based on ':-^{width}}")
    for row in primary.groupBy("weight_basis").count().orderBy(F.desc("count")).collect():
        print(f"  {row['weight_basis']:<36}{row['count']:>{width - 38},}")
    print("=" * width)


def report_census(census):
    """Print SA2 coverage and nulls per column."""
    width = 58
    n_sa2 = census.count()
    print(f"{' ABS SA2 table ':-^{width}}")
    print(f"{'SA2s':<38}{n_sa2:>{width - 38},}")
    value_cols = [c for c in census.columns if c not in ("sa2_code", "sa2_name")]
    null_counts = census.select(
        [F.sum(F.col(c).isNull().cast("int")).alias(c) for c in value_cols]
    ).first()
    for col in value_cols:
        print(f"  {col:<46}{null_counts[col]:>{width - 48},}")
    print("=" * width)


def main():
    spark = SparkSession.builder.config("spark.sql.session.timeZone", "Etc/UTC").appName("bnpl_external_abs_sa2").getOrCreate()

    EXTERNAL_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    census_zip = download(CENSUS_SA2_URL, EXTERNAL_DIR / "2021_GCP_SA2_for_AUS_short-header.zip")
    seifa_xlsx = download(SEIFA_SA2_URL, EXTERNAL_DIR / "SEIFA_2021_SA2_indexes.xlsx")
    mb_alloc_xlsx = download(MB_ALLOCATION_URL, EXTERNAL_DIR / "MB_2021_AUST.xlsx")
    poa_alloc_xlsx = download(POA_ALLOCATION_URL, EXTERNAL_DIR / "POA_2021_AUST.xlsx")
    mb_counts_xlsx = download(MB_COUNTS_URL, EXTERNAL_DIR / "Mesh_Block_Counts_2021.xlsx")
    sa2_shp_zip = download(SA2_BOUNDARY_URL, EXTERNAL_DIR / "SA2_2021_AUST_SHP_GDA2020.zip")

    print("Extracting G01 and G02 ...")
    census_csvs = extract_census_tables(census_zip)

    seifa_csv = sheet_to_csv(
        seifa_xlsx,
        EXTERNAL_DIR / "SEIFA_2021_SA2_indexes.csv",
        header_marker="2021 Statistical Area Level 2  (SA2) 9-Digit Code",
        columns=SEIFA_COLUMNS,
        sheets=seifa_sheets(seifa_xlsx),
    )
    mb_alloc_csv = sheet_to_csv(
        mb_alloc_xlsx,
        EXTERNAL_DIR / "MB_2021_AUST.csv",
        header_marker="MB_CODE_2021",
        columns=[
            "MB_CODE_2021", "MB_CATEGORY_2021", "CHANGE_FLAG_2021", "CHANGE_LABEL_2021",
            "SA1_CODE_2021", "SA2_CODE_2021", "SA2_NAME_2021", "SA3_CODE_2021", "SA3_NAME_2021",
            "SA4_CODE_2021", "SA4_NAME_2021", "GCCSA_CODE_2021", "GCCSA_NAME_2021",
            "STATE_CODE_2021", "STATE_NAME_2021",
        ],
    )
    poa_alloc_csv = sheet_to_csv(
        poa_alloc_xlsx,
        EXTERNAL_DIR / "POA_2021_AUST.csv",
        header_marker="MB_CODE_2021",
        columns=["MB_CODE_2021", "POA_CODE_2021", "POA_NAME_2021"],
    )
    mb_counts_csv = sheet_to_csv(
        mb_counts_xlsx,
        EXTERNAL_DIR / "Mesh_Block_Counts_2021.csv",
        header_marker="MB_CODE_2021",
        columns=["MB_CODE_2021", "MB_CATEGORY_NAME_2021", "AREA_ALBERS_SQKM", "Dwelling", "Person"],
        sheets=mb_counts_sheets(mb_counts_xlsx),
    )

    print("Building SA2 Census + SEIFA table ...")
    census = build_census_table(spark, census_csvs[G01_CSV], census_csvs[G02_CSV], seifa_csv)
    report_census(census)
    census.write.mode("overwrite").parquet(str(OUTPUT_DIR / "abs_census_sa2"))
    print(f"Written to {OUTPUT_DIR / 'abs_census_sa2'}")

    print("Building postcode -> SA2 correspondence ...")
    correspondence, primary = build_correspondence(
        spark, mb_alloc_csv, poa_alloc_csv, mb_counts_csv
    )
    correspondence.cache()
    primary.cache()
    report_correspondence(correspondence, primary)
    correspondence.write.mode("overwrite").parquet(str(OUTPUT_DIR / "abs_poa_sa2_correspondence"))
    primary.write.mode("overwrite").parquet(str(OUTPUT_DIR / "abs_poa_sa2_primary"))
    print(f"Written to {OUTPUT_DIR / 'abs_poa_sa2_correspondence'} and {OUTPUT_DIR / 'abs_poa_sa2_primary'}")

    write_boundaries(sa2_shp_zip, EXTERNAL_DIR / "SA2_2021_AUST_simplified.geojson")

    print("Done.")


if __name__ == "__main__":
    main()
